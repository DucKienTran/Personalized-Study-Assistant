from docx import Document

from app.services.document.parser import DocumentParserService


def test_parser_reads_pasted_text_as_single_page(tmp_path):
    text_file = tmp_path / "pasted-text.txt"
    text_file.write_text("First paragraph.\n\nSecond paragraph.", encoding="utf-8")

    parsed = DocumentParserService()._parse_sync(str(text_file))

    assert parsed.total_pages == 1
    assert parsed.markdown == "<!--page:1-->\nFirst paragraph.\n\nSecond paragraph."


def test_parser_preserves_markdown_headings(tmp_path):
    markdown_file = tmp_path / "notes.md"
    markdown_file.write_text("# Topic\n\n## Detail\n\nContent", encoding="utf-8")

    parsed = DocumentParserService()._parse_sync(str(markdown_file))

    assert parsed.total_pages == 1
    assert parsed.markdown == "<!--page:1-->\n# Topic\n\n## Detail\n\nContent"


def test_docx_parser_preserves_heading_levels_for_viewer_toc(tmp_path):
    docx_file = tmp_path / "structured.docx"
    document = Document()
    document.add_heading("Main topic", level=1)
    document.add_paragraph("Introduction")
    document.add_heading("Nested topic", level=2)
    document.add_paragraph("Details")
    document.save(docx_file)

    parsed = DocumentParserService()._parse_sync(str(docx_file))

    assert parsed.total_pages == 1
    assert "# Main topic" in parsed.markdown
    assert "## Nested topic" in parsed.markdown
